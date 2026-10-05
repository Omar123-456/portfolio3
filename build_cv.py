"""Generate Omar_Elsharoud_CV.docx from the portfolio data in app.py."""
import os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app import portfolio as p

d = Document()
sec = d.sections[0]
sec.left_margin = sec.right_margin = Cm(1.9)
sec.top_margin = sec.bottom_margin = Cm(1.6)
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(0)

def para(text="", bold=False, italic=False, size=None, align=None, after=0, before=0, color=None):
    q = d.add_paragraph(); r = q.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor(*color)
    if align: q.alignment = align
    q.paragraph_format.space_after = Pt(after); q.paragraph_format.space_before = Pt(before)
    return q

def heading(t):
    q = para(t.upper(), bold=True, size=12, before=9, after=3)
    pPr = q._p.get_or_add_pPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "444444")): bt.set(qn("w:" + k), v)
    b.append(bt); pPr.append(b)

def bullet(t, label=None):
    q = d.add_paragraph(style="List Bullet"); q.paragraph_format.space_after = Pt(1)
    if label: q.add_run(label + " ").bold = True
    q.add_run(t)

def sub(title, meta=None, before=5):
    q = para(before=before)
    q.add_run(title).bold = True
    if meta: q.add_run("  |  " + meta).italic = True
    q.paragraph_format.keep_with_next = True

C = WD_ALIGN_PARAGRAPH.CENTER
para(p["name"].upper(), bold=True, size=20, align=C)
para("Software Engineer", bold=True, size=11, align=C, after=2)
hdr = para(f'{p["location"]} | {p["email"]} | {p["phone"]}', align=C, size=9.5)
PORTFOLIO_URL = os.environ.get("PORTFOLIO_URL") or (sys.argv[1] if len(sys.argv) > 1 else "")
if PORTFOLIO_URL:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rid = d.part.relate_to(PORTFOLIO_URL, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hl = OxmlElement("w:hyperlink"); hl.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    col = OxmlElement("w:color"); col.set(qn("w:val"), "0563C1"); rPr.append(col)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "19"); rPr.append(sz)
    r.append(rPr); t = OxmlElement("w:t"); t.text = "Portfolio: " + PORTFOLIO_URL; r.append(t); hl.append(r)
    hdr.add_run(" | ").font.size = Pt(9.5)
    hdr._p.append(hl)

heading("Professional Profile")
para("Software Engineer with experience delivering enterprise business solutions on Microsoft Power Platform, "
     "Dynamics 365 and Azure, alongside full-stack web applications built in PHP, Node.js and React. At OGI, "
     "contributed to annual recurring billing automation, Dynamics customer-service enhancements, a warehouse "
     "ticketing system, a contract-import integration and the AssetTrack asset management platform, taking "
     "solutions from design through UAT to production with Azure DevOps CI/CD. Independently built and "
     "maintain the Dar Ul Isra and Bayan Academy websites, including online registration and card payments. "
     "BSc (Hons) Software Engineering graduate with a strong grounding in the full SDLC, secure systems and testing.", after=2)

heading("Technical Toolkit")
bullet(", ".join(["Microsoft Power Platform (Power Apps, Power Automate, Dataverse)", "Dynamics 365 Customer Service & Plugin Development", "Model-Driven Apps", "SharePoint Online"]), "Microsoft Business Apps:")
bullet("Microsoft Azure (App Service, PostgreSQL, Key Vault), Microsoft Entra ID, Azure DevOps, CI/CD, Docker, Git/GitHub", "Cloud & DevOps:")
for c in p["skills_categorized"]:
    if c["category"].startswith("Enterprise"): continue
    bullet(", ".join(c["items"]), c["category"] + ":")
bullet("REST API Integrations, Billing & Finance Integrations, Solution Architecture, Requirements Gathering, UAT & Production Deployments, Technical Documentation & Handover", "Delivery:")

heading("Professional Experience")
ogi = p["experience"][0]
sub("Software Engineer, OGI", ogi["dates"], before=2)
for pr in ogi["projects"]:
    q = para(before=4); q.add_run(pr["title"]).bold = True; q.add_run("  |  " + pr["role"]).italic = True
    q.paragraph_format.keep_with_next = True
    para(pr["desc"], after=1)
    q = para(); q.add_run("Technologies: ").bold = True; q.add_run(", ".join(pr["tech"])); q.paragraph_format.space_after = Pt(2)

for ex in p["experience"][1:]:
    sub(ex["company"] + " - " + ex["role"], ex["dates"], before=8)
    for h in ex["highlights"]: bullet(h)

heading("Technical Projects")
order = ["Bayan Academy Website", "Dar Ul Isra Platform", "Align", "WhatsApp AI Automation SaaS", "AccomFix",
         "CampusTasker", "QuizCraft", "Mental Health Support Platform"]
byt = {x["title"]: x for x in p["projects"]}
for t in order:
    sub(t, before=4)
    for i in byt[t]["items"]: bullet(i)

heading("Education")
for e in p["education"]:
    sub(e["degree"] + ", " + e["institution"], e["dates"], before=3)
    bullet(e["notes"])

heading("Leadership & Volunteering")
for l in p["leadership"]:
    sub(l["role"], l["org"], before=3)
    for h in l["highlights"]: bullet(h)

heading("Languages & Additional Information")
bullet(", ".join(p["languages"]), "Languages:")
for a in p["additional_info"]: bullet(a)

d.save("Omar_Elsharoud_CV.docx")
